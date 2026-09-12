from io import BytesIO
from typing import Any, Dict

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models import Company, Site


def _is_active(value) -> bool:
    if value is None:
        return True
    return str(value).strip() != "غیرفعال"


def _sheet_rows(ws):
    header_row = next(ws.iter_rows(min_row=1, max_row=1))
    headers = [str(c.value).strip() if c.value is not None else "" for c in header_row]
    for row in ws.iter_rows(min_row=2):
        values = [c.value for c in row]
        if all(v is None for v in values):
            continue
        yield dict(zip(headers, values))


def import_companies_sites(file_bytes: bytes, db: Session) -> Dict[str, Any]:
    wb = load_workbook(BytesIO(file_bytes), data_only=True)
    result = {
        "companies": {"created": 0, "updated": 0, "errors": []},
        "sites": {"created": 0, "updated": 0, "errors": []},
    }

    if "شرکت‌ها" in wb.sheetnames:
        ws = wb["شرکت‌ها"]
        for i, row in enumerate(_sheet_rows(ws), start=2):
            code = row.get("کد شرکت")
            name = row.get("نام شرکت")
            if not code or not name:
                result["companies"]["errors"].append(f"ردیف {i}: کد یا نام شرکت خالی است")
                continue
            code = str(code).strip()
            name = str(name).strip()
            active = _is_active(row.get("وضعیت"))
            obj = db.query(Company).filter(Company.code == code).first()
            if obj:
                obj.name = name
                obj.is_active = active
                result["companies"]["updated"] += 1
            else:
                db.add(Company(code=code, name=name, is_active=active))
                result["companies"]["created"] += 1
        db.commit()
    else:
        result["companies"]["errors"].append("شیت «شرکت‌ها» پیدا نشد")

    if "سایت‌ها" in wb.sheetnames:
        ws = wb["سایت‌ها"]
        for i, row in enumerate(_sheet_rows(ws), start=2):
            company_code = row.get("کد شرکت")
            site_code = row.get("کد سایت")
            name = row.get("نام سایت")
            if not company_code or not site_code or not name:
                result["sites"]["errors"].append(f"ردیف {i}: کد شرکت، کد سایت یا نام سایت خالی است")
                continue
            company_code = str(company_code).strip()
            site_code = str(site_code).strip()
            name = str(name).strip()
            company = db.query(Company).filter(Company.code == company_code).first()
            if not company:
                result["sites"]["errors"].append(f"ردیف {i}: شرکت با کد {company_code} پیدا نشد")
                continue
            address = row.get("آدرس")
            address = str(address).strip() if address else None
            active = _is_active(row.get("وضعیت"))
            obj = (
                db.query(Site)
                .filter(Site.company_id == company.id, Site.code == site_code)
                .first()
            )
            if obj:
                obj.name = name
                obj.address = address
                obj.is_active = active
                result["sites"]["updated"] += 1
            else:
                db.add(
                    Site(
                        company_id=company.id,
                        code=site_code,
                        name=name,
                        address=address,
                        is_active=active,
                    )
                )
                result["sites"]["created"] += 1
        db.commit()
    else:
        result["sites"]["errors"].append("شیت «سایت‌ها» پیدا نشد")

    return result


def import_foods(file_bytes: bytes, db: Session) -> Dict[str, Any]:
    from app.models import Food

    wb = load_workbook(BytesIO(file_bytes), data_only=True)
    result = {"foods": {"created": 0, "updated": 0, "errors": []}}

    if "غذاها" not in wb.sheetnames:
        result["foods"]["errors"].append("شیت «غذاها» پیدا نشد")
        return result

    ws = wb["غذاها"]
    for i, row in enumerate(_sheet_rows(ws), start=2):
        code = row.get("کد غذا")
        name = row.get("نام غذا")
        if not code or not name:
            result["foods"]["errors"].append(f"ردیف {i}: کد یا نام غذا خالی است")
            continue
        code = str(code).strip()
        name = str(name).strip()
        category = row.get("دسته‌بندی")
        category = str(category).strip() if category else None
        active = _is_active(row.get("وضعیت"))
        obj = db.query(Food).filter(Food.code == code).first()
        if obj:
            obj.name = name
            obj.category = category
            obj.is_active = active
            result["foods"]["updated"] += 1
        else:
            db.add(Food(code=code, name=name, category=category, is_active=active))
            result["foods"]["created"] += 1
    db.commit()
    return result


def import_employees(file_bytes: bytes, db: Session) -> Dict[str, Any]:
    from app.models import Employee

    wb = load_workbook(BytesIO(file_bytes), data_only=True)
    result = {"employees": {"created": 0, "updated": 0, "errors": []}}

    if "پرسنل" not in wb.sheetnames:
        result["employees"]["errors"].append("شیت «پرسنل» پیدا نشد")
        return result

    ws = wb["پرسنل"]
    for i, row in enumerate(_sheet_rows(ws), start=2):
        personnel_code = row.get("کد پرسنلی")
        full_name = row.get("نام و نام خانوادگی")
        company_code = row.get("کد شرکت")
        site_code = row.get("کد سایت")
        if not personnel_code or not full_name or not company_code or not site_code:
            result["employees"]["errors"].append(
                f"ردیف {i}: کد پرسنلی، نام، کد شرکت یا کد سایت خالی است"
            )
            continue
        personnel_code = str(personnel_code).strip()
        full_name = str(full_name).strip()
        company_code = str(company_code).strip()
        site_code = str(site_code).strip()

        company = db.query(Company).filter(Company.code == company_code).first()
        if not company:
            result["employees"]["errors"].append(
                f"ردیف {i}: شرکت با کد {company_code} پیدا نشد"
            )
            continue
        site = (
            db.query(Site)
            .filter(Site.company_id == company.id, Site.code == site_code)
            .first()
        )
        if not site:
            result["employees"]["errors"].append(
                f"ردیف {i}: سایت با کد {site_code} در شرکت {company_code} پیدا نشد"
            )
            continue

        mobile = row.get("موبایل")
        mobile = str(mobile).strip() if mobile else None
        role = row.get("نقش")
        role = str(role).strip() if role else "کارمند"
        active = _is_active(row.get("وضعیت"))

        if mobile:
            conflict = (
                db.query(Employee)
                .filter(Employee.mobile == mobile, Employee.personnel_code != personnel_code)
                .first()
            )
            if conflict:
                result["employees"]["errors"].append(
                    f"ردیف {i}: شماره موبایل {mobile} قبلاً برای کد پرسنلی {conflict.personnel_code} ثبت شده؛ موبایل این ردیف خالی گذاشته شد"
                )
                mobile = None

        obj = db.query(Employee).filter(Employee.personnel_code == personnel_code).first()
        if obj:
            obj.full_name = full_name
            obj.mobile = mobile
            obj.company_id = company.id
            obj.site_id = site.id
            obj.role = role
            obj.is_active = active
            result["employees"]["updated"] += 1
        else:
            db.add(
                Employee(
                    personnel_code=personnel_code,
                    full_name=full_name,
                    mobile=mobile,
                    company_id=company.id,
                    site_id=site.id,
                    role=role,
                    is_active=active,
                )
            )
            result["employees"]["created"] += 1
        db.flush()
    db.commit()
    return result


def import_welfare_managers(file_bytes: bytes, db: Session) -> Dict[str, Any]:
    from app.models import Employee, WelfareManagerAssignment

    wb = load_workbook(BytesIO(file_bytes), data_only=True)
    result = {"welfare_managers": {"created": 0, "updated": 0, "errors": []}}

    if "مسئولان_رفاهی" not in wb.sheetnames:
        result["welfare_managers"]["errors"].append("شیت «مسئولان_رفاهی» پیدا نشد")
        return result

    ws = wb["مسئولان_رفاهی"]
    for i, row in enumerate(_sheet_rows(ws), start=2):
        personnel_code = row.get("EmployeeID")
        company_code = row.get("کد شرکت")
        site_code = row.get("کد سایت")
        if not personnel_code or not company_code or not site_code:
            result["welfare_managers"]["errors"].append(
                f"ردیف {i}: کد پرسنلی، کد شرکت یا کد سایت خالی است"
            )
            continue
        personnel_code = str(personnel_code).strip()
        company_code = str(company_code).strip()
        site_code = str(site_code).strip()

        employee = (
            db.query(Employee).filter(Employee.personnel_code == personnel_code).first()
        )
        if not employee:
            result["welfare_managers"]["errors"].append(
                f"ردیف {i}: پرسنل با کد {personnel_code} پیدا نشد"
            )
            continue
        company = db.query(Company).filter(Company.code == company_code).first()
        if not company:
            result["welfare_managers"]["errors"].append(
                f"ردیف {i}: شرکت با کد {company_code} پیدا نشد"
            )
            continue
        site = (
            db.query(Site)
            .filter(Site.company_id == company.id, Site.code == site_code)
            .first()
        )
        if not site:
            result["welfare_managers"]["errors"].append(
                f"ردیف {i}: سایت با کد {site_code} در شرکت {company_code} پیدا نشد"
            )
            continue

        assignment_type = row.get("اصلی/جانشین")
        assignment_type = str(assignment_type).strip() if assignment_type else "اصلی"
        active = _is_active(row.get("وضعیت"))

        obj = (
            db.query(WelfareManagerAssignment)
            .filter(
                WelfareManagerAssignment.employee_id == employee.id,
                WelfareManagerAssignment.site_id == site.id,
            )
            .first()
        )
        if obj:
            obj.assignment_type = assignment_type
            obj.is_active = active
            result["welfare_managers"]["updated"] += 1
        else:
            db.add(
                WelfareManagerAssignment(
                    employee_id=employee.id,
                    site_id=site.id,
                    assignment_type=assignment_type,
                    is_active=active,
                )
            )
            result["welfare_managers"]["created"] += 1
        db.flush()
    db.commit()
    return result


def import_welfare_managers(file_bytes: bytes, db: Session) -> Dict[str, Any]:
    from app.models import Employee, WelfareManagerAssignment

    wb = load_workbook(BytesIO(file_bytes), data_only=True)
    result = {"welfare_managers": {"created": 0, "updated": 0, "errors": []}}

    if "مسئولان_رفاهی" not in wb.sheetnames:
        result["welfare_managers"]["errors"].append("شیت «مسئولان_رفاهی» پیدا نشد")
        return result

    ws = wb["مسئولان_رفاهی"]
    for i, row in enumerate(_sheet_rows(ws), start=2):
        personnel_code = row.get("EmployeeID")
        company_code = row.get("کد شرکت")
        site_code = row.get("کد سایت")
        if not personnel_code or not company_code or not site_code:
            result["welfare_managers"]["errors"].append(
                f"ردیف {i}: کد پرسنلی، کد شرکت یا کد سایت خالی است"
            )
            continue
        personnel_code = str(personnel_code).strip()
        company_code = str(company_code).strip()
        site_code = str(site_code).strip()

        employee = (
            db.query(Employee).filter(Employee.personnel_code == personnel_code).first()
        )
        if not employee:
            result["welfare_managers"]["errors"].append(
                f"ردیف {i}: پرسنل با کد {personnel_code} پیدا نشد"
            )
            continue
        company = db.query(Company).filter(Company.code == company_code).first()
        if not company:
            result["welfare_managers"]["errors"].append(
                f"ردیف {i}: شرکت با کد {company_code} پیدا نشد"
            )
            continue
        site = (
            db.query(Site)
            .filter(Site.company_id == company.id, Site.code == site_code)
            .first()
        )
        if not site:
            result["welfare_managers"]["errors"].append(
                f"ردیف {i}: سایت با کد {site_code} در شرکت {company_code} پیدا نشد"
            )
            continue

        assignment_type = row.get("اصلی/جانشین")
        assignment_type = str(assignment_type).strip() if assignment_type else "اصلی"
        active = _is_active(row.get("وضعیت"))

        obj = (
            db.query(WelfareManagerAssignment)
            .filter(
                WelfareManagerAssignment.employee_id == employee.id,
                WelfareManagerAssignment.site_id == site.id,
            )
            .first()
        )
        if obj:
            obj.assignment_type = assignment_type
            obj.is_active = active
            result["welfare_managers"]["updated"] += 1
        else:
            db.add(
                WelfareManagerAssignment(
                    employee_id=employee.id,
                    site_id=site.id,
                    assignment_type=assignment_type,
                    is_active=active,
                )
            )
            result["welfare_managers"]["created"] += 1
        db.flush()
    db.commit()
    return result


PERSIAN_DIGITS = "۰۱۲۳۴۵۶۷۸۹"


def _parse_jalali_date(value):
    if value is None:
        return None
    import jdatetime

    text = str(value).strip()
    for i, d in enumerate(PERSIAN_DIGITS):
        text = text.replace(d, str(i))
    parts = text.replace("-", "/").split("/")
    if len(parts) != 3:
        return None
    try:
        y, m, d = int(parts[0]), int(parts[1]), int(parts[2])
        return jdatetime.date(y, m, d).togregorian()
    except (ValueError, TypeError):
        return None


def import_menu(file_bytes: bytes, db: Session) -> Dict[str, Any]:
    from app.models import Food, MenuEntry
    wb = load_workbook(BytesIO(file_bytes), data_only=True)
    result = {"menu": {"created": 0, "updated": 0, "errors": []}}
    if "منو" not in wb.sheetnames:
        result["menu"]["errors"].append("شیت «منو» پیدا نشد")
        return result

    ws = wb["منو"]
    food_columns = ["غذای ۱", "غذای ۲", "غذای ۳", "غذای ۴"]

    for i, row in enumerate(_sheet_rows(ws), start=2):
        date_val = _parse_jalali_date(row.get("تاریخ"))
        site_code = row.get("کد سایت")
        if not date_val or not site_code:
            result["menu"]["errors"].append(
                f"ردیف {i}: تاریخ یا کد سایت خالی است"
            )
            continue
        site_code = str(site_code).strip()

        site = db.query(Site).filter(Site.code == site_code).first()
        if not site:
            result["menu"]["errors"].append(
                f"ردیف {i}: سایت با کد {site_code} پیدا نشد"
            )
            continue

        status = row.get("وضعیت منو")
        status = str(status).strip() if status else "پیش‌نویس"

        any_food = False
        for display_order, col in enumerate(food_columns, start=1):
            food_name = row.get(col)
            if not food_name or not str(food_name).strip():
                continue
            any_food = True
            food_name = str(food_name).strip()
            food = db.query(Food).filter(Food.name == food_name).first()
            if not food:
                result["menu"]["errors"].append(
                    f"ردیف {i}: غذا با نام «{food_name}» پیدا نشد ({col})"
                )
                continue

            obj = (
                db.query(MenuEntry)
                .filter(
                    MenuEntry.date == date_val,
                    MenuEntry.site_id == site.id,
                    MenuEntry.food_id == food.id,
                )
                .first()
            )
            if obj:
                obj.display_order = display_order
                obj.status = status
                obj.is_selectable = True
                result["menu"]["updated"] += 1
            else:
                db.add(
                    MenuEntry(
                        date=date_val,
                        site_id=site.id,
                        food_id=food.id,
                        display_order=display_order,
                        status=status,
                        is_selectable=True,
                    )
                )
                result["menu"]["created"] += 1

        if not any_food:
            result["menu"]["errors"].append(f"ردیف {i}: هیچ غذایی وارد نشده")

    db.commit()
    return result

def import_settings(file_bytes: bytes, db: Session) -> Dict[str, Any]:
    from app.models import Setting

    wb = load_workbook(BytesIO(file_bytes), data_only=True)
    result = {"settings": {"created": 0, "updated": 0, "errors": []}}

    if "تنظیمات" not in wb.sheetnames:
        result["settings"]["errors"].append("شیت «تنظیمات» پیدا نشد")
        return result

    ws = wb["تنظیمات"]
    for i, row in enumerate(_sheet_rows(ws), start=2):
        key = row.get("کلید")
        value = row.get("مقدار")
        if not key or value is None:
            result["settings"]["errors"].append(f"ردیف {i}: کلید یا مقدار خالی است")
            continue
        key = str(key).strip()
        value = str(value).strip()
        value_type = row.get("واحد/نوع")
        value_type = str(value_type).strip() if value_type else None
        description = row.get("توضیح")
        description = str(description).strip() if description else None

        obj = db.query(Setting).filter(Setting.key == key).first()
        if obj:
            obj.value = value
            obj.value_type = value_type
            obj.description = description
            result["settings"]["updated"] += 1
        else:
            db.add(
                Setting(
                    key=key,
                    value=value,
                    value_type=value_type,
                    description=description,
                )
            )
            result["settings"]["created"] += 1
        db.flush()
    db.commit()
    return result
